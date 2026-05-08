"""Tests for the deterministic layout grading system.

Unit tests cover each layer in isolation using minimal synthetic fixtures.
Integration tests (marked with @pytest.mark.integration) run against real
samples and require rendered artefacts in tmp/artefacts/.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_REPO = Path(__file__).resolve().parents[1]
_SRC = _REPO / "src"

import sys
if str(_SRC) not in sys.path:
    sys.path.insert(0, str(_SRC))


def _make_valid_ir(
    *,
    n_body_paras: int = 3,
    with_experience: bool = True,
    n_roles: int = 2,
    n_bullets: int = 4,
) -> dict:
    """Build a minimal valid IR dict for testing."""
    para_counter = 0

    def _next_pid():
        nonlocal para_counter
        para_counter += 1
        return f"para_{para_counter}"

    def _para(text: str, semantic: str = "paragraph") -> dict:
        return {"text": text, "semantic": semantic, "paragraph_profile": None, "para_id": _next_pid()}

    header_paras = [_para("John Doe", "section_heading"), _para("john@example.com")]

    summary_section = {
        "title": "Summary",
        "semantic_type": "summary",
        "heading": _para("Summary", "section_heading"),
        "body_paras": [_para(f"Body paragraph {i}") for i in range(n_body_paras)],
        "roles": [],
        "section_id": "sec_1",
    }

    sections = [summary_section]

    if with_experience:
        roles = []
        for ri in range(n_roles):
            role = {
                "header": _para(f"Engineer | Company {ri}", "role_header"),
                "header_extra": [],
                "meta_lines": [_para(f"2020–202{ri}", "role_meta")],
                "bullets": [_para(f"Bullet {i}", "bullet") for i in range(n_bullets)],
                "role_id": f"engineer-company-{ri}",
                "role_id_stable": f"role_{ri + 1}",
            }
            roles.append(role)

        exp_section = {
            "title": "Experience",
            "semantic_type": "experience",
            "heading": _para("Experience", "section_heading"),
            "body_paras": [],
            "roles": roles,
            "section_id": "sec_2",
        }
        sections.append(exp_section)

    return {
        "source_kind": "docx",
        "header_paras": header_paras,
        "sections": sections,
        "layout": {
            "page_width_pt": 612.0,
            "page_height_pt": 792.0,
            "margin_top_pt": 36.0,
            "margin_bottom_pt": 36.0,
            "margin_left_pt": 54.0,
            "margin_right_pt": 54.0,
            "default_font_name": "Calibri",
            "default_font_size_pt": 10.0,
        },
        "label_column_fixed": False,
        "table_column_layout_fixed": False,
    }


# ---------------------------------------------------------------------------
# IR Validator
# ---------------------------------------------------------------------------

class TestIRValidator:
    def test_valid_ir_passes(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        result = validate_ir(ir)
        assert result.passed
        assert not result.hard_fail
        assert result.failures == []

    def test_empty_para_id_fails(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        ir["sections"][0]["body_paras"].append(
            {"text": "Important text", "semantic": "paragraph", "para_id": ""}
        )
        result = validate_ir(ir)
        assert not result.passed
        assert "EMPTY_PARA_ID" in result.failures

    def test_empty_para_id_not_hard_fail_by_itself(self):
        """Fix 3: EMPTY_PARA_ID is no longer unconditionally a hard-fail at the
        validator level; grader.py applies conditional logic."""
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        ir["sections"][0]["body_paras"].append(
            {"text": "Mis-parsed role header", "semantic": "paragraph", "para_id": ""}
        )
        result = validate_ir(ir)
        assert "EMPTY_PARA_ID" in result.failures
        # Validator itself must NOT hard-fail for a single empty para_id
        assert not result.hard_fail

    def test_empty_para_id_count_tracked(self):
        """Fix 3: empty_para_id_count is populated accurately."""
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        for i in range(3):
            ir["sections"][0]["body_paras"].append(
                {"text": f"Unbound text {i}", "semantic": "paragraph", "para_id": ""}
            )
        result = validate_ir(ir)
        assert result.empty_para_id_count == 3

    def test_empty_para_id_count_zero_when_no_violations(self):
        """Fix 3: empty_para_id_count is 0 for a clean IR."""
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        result = validate_ir(ir)
        assert result.empty_para_id_count == 0

    def test_empty_text_para_with_empty_id_is_ok(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        # Empty text + empty para_id is fine (empty paragraphs)
        ir["sections"][0]["body_paras"].append(
            {"text": "", "semantic": "empty", "para_id": ""}
        )
        result = validate_ir(ir)
        assert result.passed

    def test_duplicate_para_id_fails(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        # Use para_1 again (already used by header_paras[0])
        ir["sections"][0]["body_paras"].append(
            {"text": "Dup text", "semantic": "paragraph", "para_id": "para_1"}
        )
        result = validate_ir(ir)
        assert not result.passed
        assert "DUPLICATE_PARA_ID" in result.failures

    def test_experience_no_roles_fails(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir(with_experience=False)
        # Add experience section with no roles
        ir["sections"].append({
            "title": "Experience",
            "semantic_type": "experience",
            "heading": {"text": "Experience", "semantic": "section_heading", "para_id": "para_x99"},
            "body_paras": [],
            "roles": [],
            "section_id": "sec_x",
        })
        result = validate_ir(ir)
        assert not result.passed
        assert "EXPERIENCE_NO_ROLES" in result.failures

    def test_current_date_leakage_fails(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        ir["sections"][0]["body_paras"][0]["text"] = "Working since Current Date on this project"
        result = validate_ir(ir)
        assert not result.passed
        assert "CURRENT_DATE_LEAKAGE" in result.failures

    def test_empty_section_id_fails(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir()
        ir["sections"][0]["section_id"] = ""
        result = validate_ir(ir)
        assert not result.passed
        assert "EMPTY_SECTION_ID" in result.failures

    def test_role_boundary_corruption_fails(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir
        ir = _make_valid_ir(with_experience=True, n_roles=1)
        # Corrupt a bullet to have section_heading semantic
        exp_sec = next(s for s in ir["sections"] if s["semantic_type"] == "experience")
        exp_sec["roles"][0]["bullets"][0]["semantic"] = "section_heading"
        result = validate_ir(ir)
        assert not result.passed
        assert "ROLE_BOUNDARY_CORRUPTION" in result.failures

    def test_ir_score_perfect(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir, ir_score
        ir = _make_valid_ir()
        result = validate_ir(ir)
        assert ir_score(result) == 100.0

    def test_ir_score_degraded_on_failure(self):
        from tailor.eval.layout_grader.ir_validator import validate_ir, ir_score
        ir = _make_valid_ir(with_experience=False)
        ir["sections"].append({
            "title": "Experience",
            "semantic_type": "experience",
            "heading": {"text": "Experience", "semantic": "section_heading", "para_id": "px1"},
            "body_paras": [],
            "roles": [],
            "section_id": "sx1",
        })
        result = validate_ir(ir)
        assert ir_score(result) < 100.0


# ---------------------------------------------------------------------------
# DOCX Comparator
# ---------------------------------------------------------------------------

def _create_docx(path: Path, n_paragraphs: int = 15, n_tables: int = 0, has_word_cols: bool = False):
    """Create a minimal DOCX at path for testing."""
    from docx import Document

    _W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    doc = Document()
    for i in range(n_paragraphs):
        doc.add_paragraph(f"Paragraph {i} content here")
    for _ in range(n_tables):
        doc.add_table(rows=2, cols=2)

    if has_word_cols:
        body = doc.element.body
        sect_pr = body.find(f"{{{_W}}}sectPr")
        if sect_pr is not None:
            # python-docx default template already contains w:cols (for spacing).
            # Modify the existing element rather than adding a duplicate.
            cols_el = sect_pr.find(f"{{{_W}}}cols")
            if cols_el is None:
                from lxml import etree
                cols_el = etree.SubElement(sect_pr, f"{{{_W}}}cols")
            cols_el.set(f"{{{_W}}}num", "2")

    doc.save(str(path))


class TestDocxComparator:
    def test_identical_structure_high_score(self, tmp_path):
        from tailor.eval.layout_grader.docx_comparator import compare_docx_structure
        orig = tmp_path / "orig.docx"
        gen = tmp_path / "gen.docx"
        _create_docx(orig, 20)
        _create_docx(gen, 20)
        result = compare_docx_structure(str(orig), str(gen))
        assert result.score >= 90
        assert result.paragraph_ratio == pytest.approx(1.0)
        assert not result.columns_lost
        assert not result.renderer_fallback

    def test_paragraph_collapse_penalized(self, tmp_path):
        from tailor.eval.layout_grader.docx_comparator import compare_docx_structure
        orig = tmp_path / "orig.docx"
        gen = tmp_path / "gen.docx"
        _create_docx(orig, 20)
        _create_docx(gen, 10)   # 50% collapse
        result = compare_docx_structure(str(orig), str(gen))
        assert result.score < 80
        assert result.paragraph_ratio == pytest.approx(0.5, abs=0.05)
        assert result.renderer_fallback

    def test_paragraph_expansion_penalized(self, tmp_path):
        from tailor.eval.layout_grader.docx_comparator import compare_docx_structure
        orig = tmp_path / "orig.docx"
        gen = tmp_path / "gen.docx"
        _create_docx(orig, 10)
        _create_docx(gen, 20)   # 2x expansion
        result = compare_docx_structure(str(orig), str(gen))
        assert result.score < 95
        assert result.paragraph_ratio > 1.0

    def test_table_loss_penalized(self, tmp_path):
        from tailor.eval.layout_grader.docx_comparator import compare_docx_structure
        orig = tmp_path / "orig.docx"
        gen = tmp_path / "gen.docx"
        _create_docx(orig, 10, n_tables=4)
        _create_docx(gen, 10, n_tables=0)
        result = compare_docx_structure(str(orig), str(gen))
        assert result.score < 80

    def test_word_columns_loss_is_hard_fail_trigger(self, tmp_path):
        from tailor.eval.layout_grader.docx_comparator import compare_docx_structure
        orig = tmp_path / "orig.docx"
        gen = tmp_path / "gen.docx"
        _create_docx(orig, 20, has_word_cols=True)
        _create_docx(gen, 20, has_word_cols=False)
        result = compare_docx_structure(str(orig), str(gen))
        assert result.columns_lost
        assert result.has_word_columns_original
        assert not result.has_word_columns_generated
        assert result.score <= 60  # substantial penalty (-40 from 100)

    def test_no_columns_both_docs_ok(self, tmp_path):
        from tailor.eval.layout_grader.docx_comparator import compare_docx_structure
        orig = tmp_path / "orig.docx"
        gen = tmp_path / "gen.docx"
        _create_docx(orig, 20, has_word_cols=False)
        _create_docx(gen, 20, has_word_cols=False)
        result = compare_docx_structure(str(orig), str(gen))
        assert not result.columns_lost


# ---------------------------------------------------------------------------
# Content Injection
# ---------------------------------------------------------------------------

class TestContentInjection:
    def _make_docx(self, path: "Path", text: str):
        from docx import Document
        doc = Document()
        for line in text.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
        doc.save(str(path))

    def test_injection_worked_high_score(self, tmp_path):
        """When rendered IR has LLM content, score should be high."""
        from tailor.eval.layout_grader.content_injection import check_content_injection

        llm_text = (
            "Gabriel Mitchell cloud architect kubernetes terraform prometheus grafana "
            "microservices containerization deployment pipeline infrastructure "
            "observability reliability scalability platform engineering"
        )
        ir_dict = _make_valid_ir()
        # Inject LLM-specific tokens into the IR text
        ir_dict["sections"][0]["body_paras"][0]["text"] = (
            "Cloud architect with kubernetes terraform prometheus grafana "
            "microservices containerization deployment pipeline infrastructure"
        )

        self._make_docx(
            tmp_path / "template.docx",
            "John Smith legacy java enterprise spring boot monolith mainframe COBOL",
        )

        result = check_content_injection(
            llm_text, ir_dict, str(tmp_path / "template.docx")
        )
        assert result.score >= 70
        assert not result.hard_fail
        assert result.sim_llm > result.sim_template

    def test_no_injection_low_score(self, tmp_path):
        """When rendered IR matches template (not LLM), score should be low."""
        from tailor.eval.layout_grader.content_injection import check_content_injection

        llm_text = (
            "Gabriel Mitchell cloud kubernetes terraform prometheus grafana microservices "
            "containerization pipeline infrastructure observability reliability scalability"
        )
        ir_dict = _make_valid_ir()
        # IR text looks like the template, not the LLM text
        template_body = (
            "John Smith legacy java enterprise spring boot monolith mainframe traditional"
        )
        ir_dict["sections"][0]["body_paras"][0]["text"] = template_body

        self._make_docx(tmp_path / "template.docx", template_body)

        result = check_content_injection(
            llm_text, ir_dict, str(tmp_path / "template.docx")
        )
        assert result.score <= 75
        assert result.sim_template >= result.sim_llm

    def test_lorem_ipsum_hard_fail(self, tmp_path):
        """Lorem ipsum in rendered IR should trigger hard fail."""
        from tailor.eval.layout_grader.content_injection import check_content_injection

        ir_dict = _make_valid_ir()
        ir_dict["sections"][0]["body_paras"][0]["text"] = (
            "Lorem ipsum dolor sit amet consectetur adipiscing elit"
        )
        self._make_docx(tmp_path / "template.docx", "Original template content here")

        result = check_content_injection("Any LLM text here", ir_dict, str(tmp_path / "template.docx"))
        assert result.hard_fail
        assert result.score == 0.0
        assert any("lorem" in e.lower() for e in result.evidence)

    def test_default_score_when_no_gen_json(self, tmp_path):
        """grade_sample without gen_json_path uses content injection default."""
        from tailor.eval.layout_grader.grader import grade_sample

        ir = _make_valid_ir()
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")
        _create_docx(tmp_path / "template.docx", 20)
        _create_docx(tmp_path / "generated.docx", 20)

        grade = grade_sample(
            sample_id="test_no_gen",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            # gen_json_path omitted
            pdf_method="local",
        )
        # Default content injection score is 80
        assert grade.metrics["content_injection_score"] == 80.0


# ---------------------------------------------------------------------------
# PDF Scorer — pure sub-function tests (no real PDFs needed)
# ---------------------------------------------------------------------------

class TestPDFScorerSubFunctions:
    def test_page_count_no_growth(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_page_count_score
        score, fail, _ = _compute_page_count_score(2, 2)
        assert score == 100.0
        assert not fail

    def test_page_count_one_page_growth_ok(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_page_count_score
        score, fail, _ = _compute_page_count_score(2, 3)
        assert score == 80.0  # +1 page is a soft penalty, not a pass
        assert not fail

    def test_page_count_two_page_growth_warning(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_page_count_score
        score, fail, _ = _compute_page_count_score(2, 4)
        assert score == 70.0
        assert not fail

    def test_page_count_overflow_hard_fail(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_page_count_score
        score, fail, ev = _compute_page_count_score(2, 5)
        assert score == 0.0
        assert fail
        assert any("HARD FAIL" in e for e in ev)

    def test_density_score_healthy(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_density_score_from_ir
        ir = _make_valid_ir(n_roles=3, n_bullets=5)
        score, hard_fail, evidence = _compute_density_score_from_ir(ir)
        assert score == 100.0
        assert not hard_fail
        assert not evidence

    def test_density_score_empty_roles_penalized(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_density_score_from_ir
        ir = _make_valid_ir(n_roles=4, n_bullets=0)
        score, hard_fail, evidence = _compute_density_score_from_ir(ir)
        assert score < 70.0
        assert hard_fail  # all 4 roles empty → hard fail
        assert any("bullet" in e for e in evidence)

    def test_density_score_overflow_penalized(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_density_score_from_ir
        ir = _make_valid_ir(n_roles=2, n_bullets=15)  # > 12 threshold
        score, hard_fail, evidence = _compute_density_score_from_ir(ir)
        assert score < 100.0
        assert not hard_fail  # overflow is not a hard fail
        assert any("overflow" in e for e in evidence)

    def test_region_layout_preserved(self):
        from tailor.eval.layout_grader.pdf_scorer import _score_region_layout
        from tailor.eval.models import DocumentFeatures, ExtractedDoc

        def _mock_doc(headings, col_count=1):
            return ExtractedDoc(
                path="mock.pdf",
                pages=[],
                features=DocumentFeatures(
                    page_count=2,
                    column_count_estimate=col_count,
                    column_confidence="none",
                    dominant_left_margins=[72.0],
                    dominant_line_gap=14.0,
                    dominant_section_gap=28.0,
                    dominant_body_left_x=72.0,
                ),
                headings=headings,
            )

        orig = _mock_doc(["Experience", "Education", "Skills"])
        gen = _mock_doc(["Experience", "Education", "Skills"])
        score, hard_fail, ev = _score_region_layout(orig, gen)
        assert score > 80
        assert not hard_fail

    def test_region_layout_column_loss_hard_fail(self):
        from tailor.eval.layout_grader.pdf_scorer import _score_region_layout
        from tailor.eval.models import DocumentFeatures, ExtractedDoc

        def _mock_doc(headings, col_count):
            return ExtractedDoc(
                path="mock.pdf",
                pages=[],
                features=DocumentFeatures(2, col_count, "none", [], 0, 0, 0),
                headings=headings,
            )

        orig = _mock_doc(["Experience", "Education"], col_count=2)
        gen = _mock_doc(["Experience", "Education"], col_count=1)
        score, hard_fail, ev = _score_region_layout(orig, gen)
        assert hard_fail
        assert any("Column" in e for e in ev)

    def test_blank_page_detection_none(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_blank_page_score
        from tailor.eval.models import DocumentFeatures, ExtractedDoc, LineModel, PageModel

        def _line(text):
            return LineModel(
                text=text, bbox=(72, 100, 500, 114),
                spans=[], left_x=72, right_x=500, baseline_y=100,
            )

        # Need enough lines (≥ _BLANK_LINE_THRESHOLD=4) to avoid blank-page detection
        pages = [
            PageModel(
                page_number=1, width=612, height=792,
                lines=[
                    _line("Professional Summary section heading"),
                    _line("Experienced software engineer with ten years background"),
                    _line("Experience section with multiple roles and projects"),
                    _line("Education and additional skills and certifications"),
                    _line("Contact information and references available"),
                ],
                blocks=[],
            )
        ]
        doc = ExtractedDoc(
            path="mock.pdf", pages=pages,
            features=DocumentFeatures(1, 1, "none", [], 0, 0, 0),
        )
        score, fail, _, blank = _compute_blank_page_score(doc)
        assert score == 100.0
        assert not fail
        assert blank == []

    def test_blank_middle_page_hard_fail(self):
        from tailor.eval.layout_grader.pdf_scorer import _compute_blank_page_score
        from tailor.eval.models import DocumentFeatures, ExtractedDoc, LineModel, PageModel

        def _line(text):
            return LineModel(
                text=text, bbox=(72, 100, 500, 114),
                spans=[], left_x=72, right_x=500, baseline_y=100,
            )

        pages = [
            PageModel(page_number=1, width=612, height=792,
                      lines=[_line("Page 1 content"), _line("More text here")], blocks=[]),
            # Page 2 is blank (no lines)
            PageModel(page_number=2, width=612, height=792, lines=[], blocks=[]),
            PageModel(page_number=3, width=612, height=792,
                      lines=[_line("Page 3 content"), _line("More here too")], blocks=[]),
        ]
        doc = ExtractedDoc(
            path="mock.pdf", pages=pages,
            features=DocumentFeatures(3, 1, "none", [], 0, 0, 0),
        )
        score, fail, ev, blank = _compute_blank_page_score(doc)
        assert score == 0.0
        assert fail
        assert 2 in blank
        assert any("HARD FAIL" in e for e in ev)


# ---------------------------------------------------------------------------
# Grader — composite scoring
# ---------------------------------------------------------------------------

class TestGraderComposite:
    def test_hard_fail_caps_score(self, tmp_path):
        """A hard-fail IR must cap the composite at 30."""
        from tailor.eval.layout_grader.grader import grade_sample

        # Write a broken IR (experience with no roles)
        ir = _make_valid_ir(with_experience=False)
        ir["sections"].append({
            "title": "Experience",
            "semantic_type": "experience",
            "heading": {"text": "Experience", "semantic": "section_heading", "para_id": "px1"},
            "body_paras": [],
            "roles": [],
            "section_id": "sx1",
        })
        ir_path = tmp_path / "broken_IR.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")

        _create_docx(tmp_path / "template.docx", 20)
        _create_docx(tmp_path / "generated.docx", 20)

        grade = grade_sample(
            sample_id="test_broken",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",  # no LibreOffice needed in unit tests
        )
        assert grade.hard_fail
        assert grade.composite_score <= 30.0
        assert grade.status == "hard_fail"
        assert "EXPERIENCE_NO_ROLES" in grade.hard_fail_reasons

    def test_perfect_structure_high_score(self, tmp_path):
        """Valid IR + identical DOCX structures should score well."""
        from tailor.eval.layout_grader.grader import grade_sample

        ir = _make_valid_ir(n_roles=3, n_bullets=4)
        ir_path = tmp_path / "good_IR.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")

        _create_docx(tmp_path / "template.docx", 25)
        _create_docx(tmp_path / "generated.docx", 25)

        grade = grade_sample(
            sample_id="test_perfect",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        # Even without PDFs (may fail on local conversion) IR + DOCX carry 40%
        assert not grade.hard_fail
        # With perfect IR (20%) and perfect DOCX (20%) = 40% of max, plus PDF defaults
        assert grade.composite_score >= 60

    def test_column_loss_hard_fail(self, tmp_path):
        """Losing w:cols → HARD FAIL regardless of other scores."""
        from tailor.eval.layout_grader.grader import grade_sample

        ir = _make_valid_ir()
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")

        _create_docx(tmp_path / "template.docx", 20, has_word_cols=True)
        _create_docx(tmp_path / "generated.docx", 20, has_word_cols=False)

        grade = grade_sample(
            sample_id="test_cols",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        assert grade.hard_fail
        assert "COLUMNS_LOST" in grade.hard_fail_reasons
        assert grade.composite_score <= 30.0

    def test_paragraph_collapse_detected(self, tmp_path):
        """A severe paragraph ratio collapse should be detected as renderer fallback."""
        from tailor.eval.layout_grader.grader import grade_sample

        ir = _make_valid_ir()
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")

        _create_docx(tmp_path / "template.docx", 40)
        _create_docx(tmp_path / "generated.docx", 10)  # 25% of original — heavy collapse

        grade = grade_sample(
            sample_id="test_collapse",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        # Paragraph ratio collapse → renderer_fallback detected
        assert grade.facts["paragraph_ratio"] < 0.5
        assert "C_OVERFLOW_FIT" in grade.failure_classes

    def test_density_overflow_penalized(self, tmp_path):
        """IR with over-dense roles (>12 bullets) reduces density score."""
        from tailor.eval.layout_grader.grader import grade_sample

        ir = _make_valid_ir(n_roles=2, n_bullets=15)
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")

        _create_docx(tmp_path / "template.docx", 20)
        _create_docx(tmp_path / "generated.docx", 20)

        grade = grade_sample(
            sample_id="test_density",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        # Density score should be < 100
        assert grade.metrics["density_score"] < 100.0


class TestConditionalEmptyParaIdHardFail:
    """Fix 3: EMPTY_PARA_ID is a conditional hard-fail (not unconditional).

    Single empty para_id + perfect content injection + perfect region → WARNING.
    Multiple empty para_ids (>2) → HARD FAIL.
    """

    def _ir_with_empty_para_ids(self, count: int) -> dict:
        """Build IR with `count` non-empty paras that have no para_id."""
        ir = _make_valid_ir()
        for i in range(count):
            ir["sections"][0]["body_paras"].append(
                {"text": f"Mis-parsed header {i}", "semantic": "paragraph", "para_id": ""}
            )
        return ir

    def test_single_empty_para_id_is_warning_not_hard_fail(self, tmp_path):
        """Fix 3: 1 empty para_id with no other degradation → not a hard-fail."""
        ir = self._ir_with_empty_para_ids(1)
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")
        _create_docx(tmp_path / "template.docx", 25)
        _create_docx(tmp_path / "generated.docx", 25)

        from tailor.eval.layout_grader.grader import grade_sample
        grade = grade_sample(
            sample_id="test_single_epi",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        # Single empty para_id + no ci_score degradation → must NOT be hard_fail
        assert not grade.hard_fail, (
            f"Single EMPTY_PARA_ID should not hard-fail; got status={grade.status!r}, "
            f"reasons={grade.hard_fail_reasons}"
        )
        assert grade.status != "hard_fail"

    def test_many_empty_para_ids_is_hard_fail(self, tmp_path):
        """Fix 3: 3+ empty para_ids → HARD FAIL (count > 2 threshold)."""
        ir = self._ir_with_empty_para_ids(3)
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")
        _create_docx(tmp_path / "template.docx", 25)
        _create_docx(tmp_path / "generated.docx", 25)

        from tailor.eval.layout_grader.grader import grade_sample
        grade = grade_sample(
            sample_id="test_many_epi",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        assert grade.hard_fail
        assert "EMPTY_PARA_ID" in grade.hard_fail_reasons
        assert grade.composite_score <= 30.0

    def test_experience_no_roles_still_hard_fails(self, tmp_path):
        """Fix 3: EXPERIENCE_NO_ROLES is still always a hard-fail (unchanged)."""
        ir = _make_valid_ir(with_experience=False)
        ir["sections"].append({
            "title": "Experience",
            "semantic_type": "experience",
            "heading": {"text": "Experience", "semantic": "section_heading", "para_id": "pz1"},
            "body_paras": [],
            "roles": [],
            "section_id": "sz1",
        })
        ir_path = tmp_path / "ir.json"
        ir_path.write_text(json.dumps(ir), encoding="utf-8")
        _create_docx(tmp_path / "template.docx", 25)
        _create_docx(tmp_path / "generated.docx", 25)

        from tailor.eval.layout_grader.grader import grade_sample
        grade = grade_sample(
            sample_id="test_no_roles",
            template_docx_path=str(tmp_path / "template.docx"),
            generated_docx_path=str(tmp_path / "generated.docx"),
            template_pdf_path=str(tmp_path / "template.pdf"),
            generated_pdf_path=str(tmp_path / "generated.pdf"),
            ir_path=str(ir_path),
            pdf_method="local",
        )
        assert grade.hard_fail
        assert "EXPERIENCE_NO_ROLES" in grade.hard_fail_reasons


# ---------------------------------------------------------------------------
# Report helpers
# ---------------------------------------------------------------------------

class TestReport:
    def _make_grade(self, sample_id: str, score: float, hard_fail: bool = False) -> "SampleGrade":
        from tailor.eval.layout_grader.grader import SampleGrade
        status = "hard_fail" if hard_fail else ("pass" if score >= 75 else ("warning" if score >= 60 else "fail"))
        return SampleGrade(
            sample_id=sample_id,
            status=status,
            composite_score=score,
            hard_fail=hard_fail,
            hard_fail_reasons=["COLUMNS_LOST"] if hard_fail else [],
            metrics={},
            facts={},
            failure_classes=["F_TOPOLOGY_COLLAPSE"] if hard_fail else [],
            evidence=[],
        )

    def test_aggregate_counts(self):
        from tailor.eval.layout_grader.report import build_aggregate
        grades = [
            self._make_grade("a", 80.0),   # pass
            self._make_grade("b", 65.0),   # warning
            self._make_grade("c", 45.0),   # fail
            self._make_grade("d", 20.0, hard_fail=True),   # hard_fail
        ]
        agg = build_aggregate(grades)
        assert agg["sample_count"] == 4
        assert agg["pass_count"] == 1
        assert agg["warning_count"] == 1
        assert agg["fail_count"] == 2  # fail + hard_fail
        assert agg["hard_fail_count"] == 1
        assert agg["worst_samples"][0] == "d"

    def test_aggregate_scores(self):
        from tailor.eval.layout_grader.report import build_aggregate
        grades = [
            self._make_grade("a", 80.0),
            self._make_grade("b", 60.0),
        ]
        agg = build_aggregate(grades)
        assert agg["average_score"] == pytest.approx(70.0)
        assert agg["min_score"] == 60.0
        assert agg["max_score"] == 80.0

    def test_empty_aggregate(self):
        from tailor.eval.layout_grader.report import build_aggregate
        agg = build_aggregate([])
        assert agg["sample_count"] == 0
        assert agg["average_score"] == 0.0

    def test_summary_txt_written(self, tmp_path):
        from tailor.eval.layout_grader.report import build_aggregate, write_summary_txt
        grades = [
            self._make_grade("sample_1", 80.0),
            self._make_grade("sample_2", 45.0),
        ]
        agg = build_aggregate(grades)
        out = tmp_path / "summary.txt"
        write_summary_txt(grades, agg, out)
        content = out.read_text(encoding="utf-8")
        assert "LAYOUT GRADING SUMMARY" in content
        assert "sample_1" in content
        assert "sample_2" in content

    def test_baseline_regression_detected(self):
        from tailor.eval.layout_grader.report import compare_baseline
        from tailor.eval.layout_grader.grader import SampleGrade
        import tempfile, json as _json

        grades = [self._make_grade("s1", 60.0), self._make_grade("s2", 80.0)]
        baseline_data = {
            "grades": [
                {"sample_id": "s1", "composite_score": 75.0},
                {"sample_id": "s2", "composite_score": 82.0},
            ]
        }
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False) as f:
            _json.dump(baseline_data, f)
            baseline_path = f.name

        regressions = compare_baseline(grades, baseline_path)
        # s1 dropped 15 points → should appear; s2 dropped 2 → below threshold
        assert len(regressions) == 1
        assert regressions[0]["sample_id"] == "s1"
        assert regressions[0]["regression_class"] == "SEVERE_REGRESSION"


# ---------------------------------------------------------------------------
# Integration tests — require real rendered artefacts
# ---------------------------------------------------------------------------

_REND_DOCX_DIR = _REPO / "tmp" / "artefacts" / "rendering" / "docx"
_IR_DIR = _REPO / "tmp" / "artefacts" / "ir" / "docx"
_RES_DOCX_DIR = _REPO / "tests" / "samples" / "resume" / "docx"

_SAMPLE_1_STEM = "1-Leonid_Verman_Resume_Template"
_SAMPLE_31_STEM = "31-Software-Engineer-Editable-Resume-Template-Download-in-docx-7"


def _sample_available(stem: str) -> bool:
    return (
        (_REND_DOCX_DIR / (stem + ".docx")).exists()
        and (_IR_DIR / (stem + "_IR.json")).exists()
        and (_RES_DOCX_DIR / (stem + ".docx")).exists()
    )


@pytest.mark.integration
@pytest.mark.skipif(
    not _sample_available(_SAMPLE_1_STEM),
    reason="Rendered artefacts for sample 1 not found",
)
def test_real_sample_1_grades_successfully(tmp_path):
    """Sample 1 (Leonid Verman personal template) should pass or warn."""
    from tailor.eval.layout_grader.grader import grade_sample

    grade = grade_sample(
        sample_id=_SAMPLE_1_STEM,
        template_docx_path=str(_RES_DOCX_DIR / (_SAMPLE_1_STEM + ".docx")),
        generated_docx_path=str(_REND_DOCX_DIR / (_SAMPLE_1_STEM + ".docx")),
        template_pdf_path=str(tmp_path / "template.pdf"),
        generated_pdf_path=str(tmp_path / "generated.pdf"),
        ir_path=str(_IR_DIR / (_SAMPLE_1_STEM + "_IR.json")),
        pdf_method="local",
    )

    # Should at minimum produce a valid grade object
    assert grade.sample_id
    assert grade.status in ("pass", "warning", "fail", "hard_fail")
    assert 0.0 <= grade.composite_score <= 100.0
    assert isinstance(grade.metrics, dict)
    assert isinstance(grade.facts, dict)
    print(f"\n  [integration] Sample 1 score={grade.composite_score:.1f} status={grade.status}")
    for ev in grade.evidence:
        print(f"    {ev}")


@pytest.mark.integration
@pytest.mark.skipif(
    not _sample_available(_SAMPLE_31_STEM),
    reason="Rendered artefacts for sample 31 not found",
)
def test_real_sample_31_grades_successfully(tmp_path):
    """Sample 31 should produce a valid grade (known complex template)."""
    from tailor.eval.layout_grader.grader import grade_sample

    grade = grade_sample(
        sample_id=_SAMPLE_31_STEM,
        template_docx_path=str(_RES_DOCX_DIR / (_SAMPLE_31_STEM + ".docx")),
        generated_docx_path=str(_REND_DOCX_DIR / (_SAMPLE_31_STEM + ".docx")),
        template_pdf_path=str(tmp_path / "template.pdf"),
        generated_pdf_path=str(tmp_path / "generated.pdf"),
        ir_path=str(_IR_DIR / (_SAMPLE_31_STEM + "_IR.json")),
        pdf_method="local",
    )

    assert grade.sample_id
    assert grade.status in ("pass", "warning", "fail", "hard_fail")
    assert 0.0 <= grade.composite_score <= 100.0
    print(f"\n  [integration] Sample 31 score={grade.composite_score:.1f} status={grade.status}")
    for ev in grade.evidence:
        print(f"    {ev}")
