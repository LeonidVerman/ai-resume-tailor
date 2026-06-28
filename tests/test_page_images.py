"""Tests for PageImageBlock serialization (Fix C) and full-page image sampling (Fix A).

Fix A: _extract_page_images uses 5-point sampling for full-page images so templates
       with white centers but colored edges/corners are not dropped.
Fix C: PageImageBlock.to_dict/from_dict round-trips image bytes as base64; wired into
       ResumeDocument.to_dict/from_dict so page_images survive DB/IR round-trips.
"""
from __future__ import annotations

import io
from pathlib import Path

import pytest

from tailor.compiler.models import PageImageBlock, ResumeDocument, LayoutProfile, ParaModel, ParaStyle

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_HERE = Path(__file__).parent
_PDF_DIR = _HERE / "samples" / "resume" / "pfd"
_SAMPLE_6  = _PDF_DIR / "6-Template1.pdf"
_SAMPLE_8  = _PDF_DIR / "8-Template3.pdf"
_SAMPLE_9  = _PDF_DIR / "9-Template4.pdf"
_SAMPLE_12 = _PDF_DIR / "12-Nurse-template2.pdf"
_SAMPLE_18 = _PDF_DIR / "18-Project-Engineer-Editable-Resume-Template-Download-in-docx.pdf"
_SAMPLE_39 = _PDF_DIR / "39-backend-developer-1606703830.pdf"


def _minimal_layout() -> LayoutProfile:
    return LayoutProfile(
        page_width_pt=612.0,
        page_height_pt=792.0,
        margin_top_pt=72.0,
        margin_bottom_pt=72.0,
        margin_left_pt=72.0,
        margin_right_pt=72.0,
        default_font_name="Calibri",
        default_font_size_pt=11.0,
    )


def _minimal_doc(page_images: list[PageImageBlock] | None = None) -> ResumeDocument:
    return ResumeDocument(
        header_paras=[],
        sections=[],
        layout=_minimal_layout(),
        all_paras=[],
        source_kind="pdf",
        page_images=page_images or [],
    )


def _solid_png(r: int, g: int, b: int, w: int = 4, h: int = 4) -> bytes:
    """Return a minimal PNG of a solid color using only stdlib."""
    import struct, zlib

    def png_chunk(name: bytes, data: bytes) -> bytes:
        c = zlib.crc32(name + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + name + data + struct.pack(">I", c)

    raw_rows = b""
    for _ in range(h):
        raw_rows += b"\x00" + bytes([r, g, b] * w)

    compressed = zlib.compress(raw_rows)
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + png_chunk(b"IHDR", ihdr)
        + png_chunk(b"IDAT", compressed)
        + png_chunk(b"IEND", b"")
    )


# ---------------------------------------------------------------------------
# Fix C: PageImageBlock round-trip
# ---------------------------------------------------------------------------

class TestPageImageBlockRoundTrip:

    def test_basic_fields_preserved(self):
        png = _solid_png(10, 20, 30)
        img = PageImageBlock(
            image_bytes=png,
            x_pt=12.5,
            y_pt=34.0,
            width_pt=100.0,
            height_pt=200.0,
            category="full_page_bg",
            page_index=1,
        )
        d = img.to_dict()
        restored = PageImageBlock.from_dict(d)

        assert restored.image_bytes == png
        assert restored.x_pt == 12.5
        assert restored.y_pt == 34.0
        assert restored.width_pt == 100.0
        assert restored.height_pt == 200.0
        assert restored.category == "full_page_bg"
        assert restored.page_index == 1

    def test_category_default_on_missing_key(self):
        png = _solid_png(0, 0, 0)
        img = PageImageBlock(image_bytes=png, x_pt=0, y_pt=0, width_pt=10, height_pt=10)
        d = img.to_dict()
        del d["category"]
        restored = PageImageBlock.from_dict(d)
        assert restored.category == "body_decor"

    def test_page_index_default_zero_on_missing_key(self):
        png = _solid_png(0, 0, 0)
        img = PageImageBlock(image_bytes=png, x_pt=0, y_pt=0, width_pt=10, height_pt=10, page_index=3)
        d = img.to_dict()
        del d["page_index"]
        restored = PageImageBlock.from_dict(d)
        assert restored.page_index == 0

    def test_all_category_values_round_trip(self):
        png = _solid_png(50, 100, 150)
        for cat in ("profile_photo", "header_footer_decor", "body_decor", "full_page_bg",
                    "sidebar_bg", "header_band", "footer_band"):
            img = PageImageBlock(image_bytes=png, x_pt=0, y_pt=0,
                                 width_pt=10, height_pt=10, category=cat)
            assert PageImageBlock.from_dict(img.to_dict()).category == cat

    def test_anchor_fields_default_none(self):
        """anchor_next_section_id and gap_to_anchor_pt default to None."""
        png = _solid_png(0, 0, 0)
        img = PageImageBlock(image_bytes=png, x_pt=0, y_pt=0, width_pt=10, height_pt=10)
        assert img.anchor_next_section_id is None
        assert img.gap_to_anchor_pt is None

    def test_anchor_fields_round_trip(self):
        """anchor_next_section_id and gap_to_anchor_pt survive to_dict/from_dict."""
        png = _solid_png(255, 100, 0)
        img = PageImageBlock(
            image_bytes=png, x_pt=50.0, y_pt=200.0, width_pt=400.0, height_pt=2.0,
            category="h_rule", page_index=0,
            anchor_next_section_id="sec_3",
            gap_to_anchor_pt=22.5,
        )
        d = img.to_dict()
        assert d["anchor_next_section_id"] == "sec_3"
        assert d["gap_to_anchor_pt"] == 22.5

        restored = PageImageBlock.from_dict(d)
        assert restored.anchor_next_section_id == "sec_3"
        assert restored.gap_to_anchor_pt == 22.5

    def test_anchor_fields_absent_when_none(self):
        """anchor fields must be absent from dict when None (backward compat)."""
        png = _solid_png(0, 0, 0)
        img = PageImageBlock(image_bytes=png, x_pt=0, y_pt=0, width_pt=10, height_pt=10)
        d = img.to_dict()
        assert "anchor_next_section_id" not in d
        assert "gap_to_anchor_pt" not in d

    def test_anchor_fields_missing_from_dict_backward_compat(self):
        """Old IR dicts without anchor keys → None fields, no KeyError."""
        png = _solid_png(0, 0, 0)
        img = PageImageBlock(image_bytes=png, x_pt=0, y_pt=0, width_pt=10, height_pt=10)
        d = img.to_dict()
        # Simulate old format (no anchor keys)
        d.pop("anchor_next_section_id", None)
        d.pop("gap_to_anchor_pt", None)
        restored = PageImageBlock.from_dict(d)
        assert restored.anchor_next_section_id is None
        assert restored.gap_to_anchor_pt is None


class TestResumeDocumentPageImagesRoundTrip:

    def test_empty_page_images_not_serialised(self):
        doc = _minimal_doc(page_images=[])
        d = doc.to_dict()
        assert "page_images" not in d

    def test_single_image_survives_round_trip(self):
        png = _solid_png(10, 20, 30)
        img = PageImageBlock(
            image_bytes=png, x_pt=5.0, y_pt=8.0,
            width_pt=612.0, height_pt=792.0,
            category="full_page_bg", page_index=0,
        )
        doc = _minimal_doc(page_images=[img])
        restored = ResumeDocument.from_dict(doc.to_dict())

        assert len(restored.page_images) == 1
        r = restored.page_images[0]
        assert r.image_bytes == png
        assert r.x_pt == 5.0
        assert r.y_pt == 8.0
        assert r.width_pt == 612.0
        assert r.height_pt == 792.0
        assert r.category == "full_page_bg"
        assert r.page_index == 0

    def test_multiple_images_order_preserved(self):
        imgs = [
            PageImageBlock(image_bytes=_solid_png(i*10, 0, 0),
                           x_pt=float(i), y_pt=0, width_pt=10, height_pt=10,
                           category="body_decor", page_index=0)
            for i in range(5)
        ]
        doc = _minimal_doc(page_images=imgs)
        restored = ResumeDocument.from_dict(doc.to_dict())
        assert len(restored.page_images) == 5
        for i, r in enumerate(restored.page_images):
            assert r.image_bytes == imgs[i].image_bytes

    def test_backward_compat_missing_page_images_key(self):
        """IR serialized before Fix C (no page_images key) → empty list, no error."""
        doc = _minimal_doc()
        d = doc.to_dict()
        assert "page_images" not in d  # old format
        restored = ResumeDocument.from_dict(d)
        assert restored.page_images == []

    def test_docx_source_unaffected(self):
        """DOCX-sourced documents always have empty page_images → key absent after to_dict."""
        doc = _minimal_doc()
        doc.source_kind = "docx"
        d = doc.to_dict()
        assert "page_images" not in d
        restored = ResumeDocument.from_dict(d)
        assert restored.page_images == []


# ---------------------------------------------------------------------------
# Section-anchored h_rule tests (unit-level, no PDF required)
# ---------------------------------------------------------------------------

class TestHRuleAnchorFieldsSerialization:
    """anchor_next_section_id / gap_to_anchor_pt serialization contract."""

    def test_round_trip_in_resume_document(self):
        """anchor fields survive ResumeDocument.to_dict/from_dict."""
        png = _solid_png(255, 100, 0)
        img = PageImageBlock(
            image_bytes=png, x_pt=50.0, y_pt=200.0, width_pt=400.0, height_pt=2.0,
            category="h_rule",
            anchor_next_section_id="sec_2",
            gap_to_anchor_pt=22.0,
        )
        doc = _minimal_doc(page_images=[img])
        restored = ResumeDocument.from_dict(doc.to_dict())
        r = restored.page_images[0]
        assert r.anchor_next_section_id == "sec_2"
        assert r.gap_to_anchor_pt == 22.0

    def test_unanchored_rule_round_trip_no_extra_keys(self):
        """Unanchored h_rule produces no anchor keys in dict; restores cleanly."""
        png = _solid_png(200, 50, 50)
        img = PageImageBlock(
            image_bytes=png, x_pt=0, y_pt=100.0, width_pt=400.0, height_pt=1.5,
            category="h_rule",
        )
        d = img.to_dict()
        assert "anchor_next_section_id" not in d
        assert "gap_to_anchor_pt" not in d
        restored = PageImageBlock.from_dict(d)
        assert restored.anchor_next_section_id is None
        assert restored.gap_to_anchor_pt is None


class TestBuildHRuleBorderMap:
    """_build_h_rule_border_map returns correct para_id → rule mapping."""

    def _make_doc_with_sections_and_rules(self):
        """Build a minimal ResumeDocument with two sections and two h_rules."""
        from tailor.compiler.models import (
            ResumeDocument, ResumeSection, ParaModel, ParaStyle,
            ParagraphProfile, LayoutProfile,
        )
        def _para(text, semantic="paragraph", para_id=""):
            pm = ParaModel(text=text, style=ParaStyle(), semantic=semantic,
                           paragraph_profile=ParagraphProfile())
            pm.para_id = para_id
            return pm

        sec1_heading = _para("OBJECTIVE", "section_heading", "para_2")
        sec2_heading = _para("EDUCATION", "section_heading", "para_5")

        sec1 = ResumeSection(title="OBJECTIVE", heading=sec1_heading,
                             semantic_type="other", section_id="sec_1")
        sec2 = ResumeSection(title="EDUCATION", heading=sec2_heading,
                             semantic_type="education", section_id="sec_2")

        png = _solid_png(255, 100, 0)
        rule1 = PageImageBlock(image_bytes=png, x_pt=0, y_pt=138.0,
                               width_pt=400.0, height_pt=2.0, category="h_rule",
                               anchor_next_section_id="sec_1", gap_to_anchor_pt=22.0)
        rule2 = PageImageBlock(image_bytes=png, x_pt=0, y_pt=234.0,
                               width_pt=400.0, height_pt=2.0, category="h_rule",
                               anchor_next_section_id="sec_2", gap_to_anchor_pt=22.0)
        unanchored = PageImageBlock(image_bytes=png, x_pt=0, y_pt=50.0,
                                    width_pt=400.0, height_pt=2.0, category="h_rule")

        layout = LayoutProfile(page_width_pt=612, page_height_pt=792,
                               margin_top_pt=72, margin_bottom_pt=72,
                               margin_left_pt=72, margin_right_pt=72,
                               default_font_name="Calibri", default_font_size_pt=11)

        doc = ResumeDocument(
            header_paras=[],
            sections=[sec1, sec2],
            layout=layout,
            all_paras=[sec1_heading, sec2_heading],
            source_kind="pdf",
            page_images=[rule1, rule2, unanchored],
        )
        return doc, rule1, rule2

    def test_returns_para_id_keyed_map(self):
        from tailor.compiler.docx_renderer import _build_h_rule_border_map
        doc, rule1, rule2 = self._make_doc_with_sections_and_rules()
        m = _build_h_rule_border_map(doc)
        assert "para_2" in m
        assert "para_5" in m
        assert m["para_2"] is rule1
        assert m["para_5"] is rule2

    def test_unanchored_rule_not_in_map(self):
        from tailor.compiler.docx_renderer import _build_h_rule_border_map
        doc, _, _ = self._make_doc_with_sections_and_rules()
        m = _build_h_rule_border_map(doc)
        # Unanchored rule doesn't have a section entry
        assert len(m) == 2

    def test_empty_when_no_page_images(self):
        from tailor.compiler.docx_renderer import _build_h_rule_border_map
        doc = _minimal_doc()
        assert _build_h_rule_border_map(doc) == {}


# ---------------------------------------------------------------------------
# Fix A: full-page image 5-point sampling
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _SAMPLE_6.exists(), reason="sample 6 PDF not in test fixtures")
class TestSample6FullPageImage:

    def test_sample6_full_page_bg_extracted(self):
        """Sample 6 has teal corners and a white center — must not be dropped."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_6, "rb") as f:
            ir = parse_pdf(f.read())

        # The decorative background must survive
        assert ir.pdf_diagnostics is not None
        assert ir.pdf_diagnostics["raster_images_raw"] >= 1, (
            "Sample 6 raster image was not extracted — center-pixel filter still active"
        )
        categories = [img.category for img in ir.page_images]
        assert "full_page_bg" in categories, (
            f"Expected full_page_bg in page_images, got: {categories}"
        )

    def test_sample6_round_trip_preserves_full_page_bg(self):
        """After to_dict/from_dict the full_page_bg image is still present."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_6, "rb") as f:
            ir = parse_pdf(f.read())

        restored = ResumeDocument.from_dict(ir.to_dict())
        categories = [img.category for img in restored.page_images]
        assert "full_page_bg" in categories


@pytest.mark.skipif(not _SAMPLE_18.exists(), reason="sample 18 PDF not in test fixtures")
class TestNearWhiteImageFiltered:
    """Near-white full-page images must remain filtered — no false positives."""

    def test_sample18_not_extracted(self):
        """Sample 18 is effectively white (min channel 250) — must stay filtered."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_18, "rb") as f:
            ir = parse_pdf(f.read())
        categories = [img.category for img in ir.page_images]
        assert "full_page_bg" not in categories, (
            f"Sample 18 near-white image should be filtered, got: {categories}"
        )


@pytest.mark.skipif(not _SAMPLE_39.exists(), reason="sample 39 PDF not in test fixtures")
class TestSample39FullPageBgRoundTrip:
    """Sample 39 already had full_page_bg extracted — verify it survives Fix C."""

    def test_full_page_bg_survives_serialization(self):
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_39, "rb") as f:
            ir = parse_pdf(f.read())

        original_cats = [img.category for img in ir.page_images]
        assert "full_page_bg" in original_cats

        restored = ResumeDocument.from_dict(ir.to_dict())
        restored_cats = [img.category for img in restored.page_images]
        assert "full_page_bg" in restored_cats

    def test_image_bytes_integrity_after_round_trip(self):
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_39, "rb") as f:
            ir = parse_pdf(f.read())

        original = {img.category: img.image_bytes for img in ir.page_images}
        restored = ResumeDocument.from_dict(ir.to_dict())
        for rimg in restored.page_images:
            assert rimg.image_bytes == original[rimg.category], (
                f"image_bytes mismatch for category={rimg.category}"
            )


# ---------------------------------------------------------------------------
# Fix 1: _extract_vector_lines called (sample 12 orange h_rules)
# Fix 2: Full-page cream background not dropped by avg>=240 filter (sample 12)
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not _SAMPLE_12.exists(), reason="sample 12 PDF not in test fixtures")
class TestSample12VectorLines:
    """Fix 1: _extract_vector_lines() must be called so orange rules appear."""

    def test_h_rule_images_extracted(self):
        """Sample 12 has orange horizontal rules — at least one h_rule must be present."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        categories = [img.category for img in ir.page_images]
        assert "h_rule" in categories, (
            f"Expected h_rule in page_images, got: {categories}"
        )

    def test_h_rule_count_ge_1(self):
        """At least one orange h_rule block must be extracted."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        h_rules = [img for img in ir.page_images if img.category == "h_rule"]
        assert len(h_rules) >= 1, (
            f"Expected >=1 h_rule images, got {len(h_rules)}"
        )

    def test_line_images_raw_diagnostic_present(self):
        """line_images_raw key must exist in pdf_diagnostics."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        assert ir.pdf_diagnostics is not None
        assert "line_images_raw" in ir.pdf_diagnostics


@pytest.mark.skipif(not _SAMPLE_12.exists(), reason="sample 12 PDF not in test fixtures")
class TestSample12CreameBackground:
    """Fix 2: Cream full-page background must not be dropped by avg>=240 filter."""

    def test_full_page_bg_extracted(self):
        """Sample 12 has a cream background (avg≈246, min<240) — must survive."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        categories = [img.category for img in ir.page_images]
        assert "full_page_bg" in categories, (
            f"Expected full_page_bg in page_images (cream bg filtered), got: {categories}"
        )

    def test_full_page_bg_survives_round_trip(self):
        """Cream full_page_bg must survive to_dict/from_dict."""
        from tailor.compiler.pdf_parser import parse_pdf
        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        restored = ResumeDocument.from_dict(ir.to_dict())
        categories = [img.category for img in restored.page_images]
        assert "full_page_bg" in categories


@pytest.mark.skipif(not _SAMPLE_12.exists(), reason="sample 12 PDF not in test fixtures")
class TestSample12HRulesAnchored:
    """Sample 12: all 4 orange h_rules (gap=32pt) must be anchored to sections."""

    def test_all_h_rules_anchored(self):
        from tailor.compiler.pdf_parser import parse_pdf

        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        h_rules = [img for img in ir.page_images if img.category == "h_rule"]
        assert len(h_rules) >= 1, "Sample 12 should have h_rule images"
        for rule in h_rules:
            assert rule.anchor_next_section_id is not None, (
                f"Sample 12 h_rule at y={rule.y_pt:.0f} should be anchored"
            )
            assert rule.gap_to_anchor_pt is not None
            assert 0 < rule.gap_to_anchor_pt <= 80.0

    def test_consistent_gap_across_rules(self):
        """All sample 12 h_rules share the same template gap (≈32pt)."""
        from tailor.compiler.pdf_parser import parse_pdf

        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        gaps = [
            img.gap_to_anchor_pt
            for img in ir.page_images
            if img.category == "h_rule" and img.gap_to_anchor_pt is not None
        ]
        assert len(gaps) >= 2, "Need at least 2 rules to test consistency"
        assert max(gaps) - min(gaps) < 2.0, (
            f"Expected consistent gap across rules, got: {gaps}"
        )

    def test_anchor_fields_survive_round_trip(self):
        """anchor fields must survive to_dict/from_dict for sample 12 rules."""
        from tailor.compiler.pdf_parser import parse_pdf

        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        restored = ResumeDocument.from_dict(ir.to_dict())
        h_rules = [img for img in restored.page_images if img.category == "h_rule"]
        for rule in h_rules:
            assert rule.anchor_next_section_id is not None, (
                "anchor_next_section_id lost in round-trip"
            )
            assert rule.gap_to_anchor_pt is not None


# ---------------------------------------------------------------------------
# Header-zone rule suppression / anchored-rule rendering (renderer guard)
# Helpers
# ---------------------------------------------------------------------------

def _rendered_floating_image_ys(ir, resume_template: str) -> "list[float]":
    """Render *ir* to a DOCX and return the y-offsets (pt) of all floating images."""
    import tempfile, os
    from tailor.compiler.docx_renderer import render_docx
    from docx import Document

    out = tempfile.mktemp(suffix=".docx")
    try:
        render_docx(ir, resume_template, out)
        d = Document(out)
        _WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        _PT = 12700
        ys = []
        for drawing in d.element.body.iter(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
        ):
            for anchor in drawing.iter(f"{{{_WP}}}anchor"):
                posV = anchor.find(f"{{{_WP}}}positionV")
                if posV is not None:
                    off = posV.find(f"{{{_WP}}}posOffset")
                    if off is not None:
                        ys.append(int(off.text) / _PT)
        return ys
    finally:
        if os.path.exists(out):
            os.unlink(out)


@pytest.mark.skipif(not _SAMPLE_8.exists(), reason="sample 8 PDF not in test fixtures")
class TestSample8HRuleAnchored:
    """Sample 8's h_rule (y=89, gap=35pt to Profile) is anchored to Profile section.

    Before: suppressed by header-zone check (y < first_sec_y).
    After: anchored to 'Profile' heading → rendered as w:pBdr/w:top border instead
    of a floating image.  The rule must NOT appear as a floating image.
    """

    def test_rule_anchored_to_profile(self):
        from tailor.compiler.pdf_parser import parse_pdf

        with open(_SAMPLE_8, "rb") as f:
            ir = parse_pdf(f.read())

        h_rules = [img for img in ir.page_images if img.category == "h_rule"]
        assert len(h_rules) >= 1, "Sample 8 should have at least one h_rule in IR"
        rule = h_rules[0]
        assert rule.anchor_next_section_id is not None, (
            f"Sample 8 h_rule at y={rule.y_pt:.0f} should be anchored (gap=35pt to Profile)"
        )
        assert rule.gap_to_anchor_pt is not None
        assert rule.gap_to_anchor_pt <= 80.0

    def test_rule_not_in_floating_images(self, tmp_path):
        from tailor.compiler.pdf_parser import parse_pdf
        from tailor.config import RESUME_TEMPLATE

        with open(_SAMPLE_8, "rb") as f:
            ir = parse_pdf(f.read())

        h_rules_in_ir = [img for img in ir.page_images if img.category == "h_rule"]
        assert len(h_rules_in_ir) >= 1, "Sample 8 should have at least one h_rule in IR"

        rendered_ys = _rendered_floating_image_ys(ir, RESUME_TEMPLATE)
        h_rule_y = h_rules_in_ir[0].y_pt
        # Use 1pt tolerance to account for int(y * 12700) / 12700 rounding
        assert not any(abs(ry - h_rule_y) < 1.0 for ry in rendered_ys), (
            f"h_rule at y={h_rule_y:.0f} must not appear as floating image (should be anchor-rendered) "
            f"(rendered floating image y-values: {[f'{y:.0f}' for y in rendered_ys]})"
        )


@pytest.mark.skipif(not _SAMPLE_9.exists(), reason="sample 9 PDF not in test fixtures")
class TestSample9AnchoredRules:
    """Sample 9: all 5 h_rules (gap≤80pt to their sections) must be anchored.

    Anchored rules are rendered as w:pBdr/w:top on section headings, NOT as
    floating images.  The old header-zone suppression logic remains but is now
    a safety net only — all body rules are diverted through the anchor path.
    """

    def test_all_body_rules_anchored(self):
        from tailor.compiler.pdf_parser import parse_pdf

        with open(_SAMPLE_9, "rb") as f:
            ir = parse_pdf(f.read())

        h_rules = sorted(
            [img for img in ir.page_images if img.category == "h_rule"],
            key=lambda img: img.y_pt,
        )
        assert len(h_rules) >= 2, "Sample 9 should have multiple h_rules"

        # Every rule that is within 80pt of a section heading must be anchored
        sec_ys = {
            sec.section_id: sec.heading.paragraph_profile.y_top_pt
            for sec in ir.sections
            if sec.heading.paragraph_profile and sec.heading.paragraph_profile.y_top_pt > 0
        }
        for rule in h_rules:
            nearest_gap = min(
                (y - rule.y_pt for y in sec_ys.values() if y > rule.y_pt),
                default=None,
            )
            if nearest_gap is not None and nearest_gap <= 80.0:
                assert rule.anchor_next_section_id is not None, (
                    f"h_rule at y={rule.y_pt:.0f} (gap={nearest_gap:.1f}pt) "
                    "should be anchored but anchor_next_section_id is None"
                )
                assert rule.gap_to_anchor_pt is not None
                assert abs(rule.gap_to_anchor_pt - nearest_gap) < 1.0

    def test_anchored_rules_not_in_floating_images(self):
        """Anchored h_rules must NOT appear as floating images in the rendered DOCX."""
        from tailor.compiler.pdf_parser import parse_pdf
        from tailor.config import RESUME_TEMPLATE

        with open(_SAMPLE_9, "rb") as f:
            ir = parse_pdf(f.read())

        anchored_ys = [
            img.y_pt for img in ir.page_images
            if img.category == "h_rule" and img.anchor_next_section_id is not None
        ]
        assert anchored_ys, "Sample 9 should have at least one anchored h_rule"

        rendered_ys = _rendered_floating_image_ys(ir, RESUME_TEMPLATE)
        for rule_y in anchored_ys:
            assert not any(abs(ry - rule_y) < 1.0 for ry in rendered_ys), (
                f"Anchored h_rule at y={rule_y:.0f} must not appear as floating image "
                f"(rendered floating ys: {[f'{y:.0f}' for y in rendered_ys]})"
            )

    def test_section_headings_have_top_border(self):
        """Section headings with anchored rules must have w:pBdr/w:top in the DOCX."""
        import tempfile, os
        from tailor.compiler.pdf_parser import parse_pdf
        from tailor.compiler.docx_renderer import render_docx
        from tailor.config import RESUME_TEMPLATE
        from docx import Document

        with open(_SAMPLE_9, "rb") as f:
            ir = parse_pdf(f.read())

        anchored_section_ids = {
            img.anchor_next_section_id
            for img in ir.page_images
            if img.category == "h_rule" and img.anchor_next_section_id is not None
        }
        assert anchored_section_ids, "Sample 9 should have anchored h_rules"

        out = tempfile.mktemp(suffix=".docx")
        try:
            render_docx(ir, RESUME_TEMPLATE, out)
            d = Document(out)
            _WM = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            borders_found = 0
            for p in d.element.body.iter(f"{{{_WM}}}p"):
                pPr = p.find(f"{{{_WM}}}pPr")
                if pPr is None:
                    continue
                pBdr = pPr.find(f"{{{_WM}}}pBdr")
                if pBdr is None:
                    continue
                top = pBdr.find(f"{{{_WM}}}top")
                if top is not None:
                    borders_found += 1
            assert borders_found >= len(anchored_section_ids), (
                f"Expected at least {len(anchored_section_ids)} w:pBdr/w:top border(s) "
                f"but found {borders_found} in rendered DOCX"
            )
        finally:
            if os.path.exists(out):
                os.unlink(out)


# ---------------------------------------------------------------------------
# Sample 12 renderer fixes
# Fix A: CONTACT INFO label-column → 3-col vMerge table
# Fix B: EDU/COMM/LEAD group table → w:tblBorders/w:top spans full width
# ---------------------------------------------------------------------------

_WM = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _render_sample_12(tmp_path):
    """Parse sample 12 PDF and render to DOCX, returning the python-docx Document."""
    from tailor.compiler.pdf_parser import parse_pdf
    from tailor.compiler.docx_renderer import render_docx
    from tailor.config import RESUME_TEMPLATE
    from docx import Document

    with open(_SAMPLE_12, "rb") as f:
        ir = parse_pdf(f.read())

    out = str(tmp_path / "sample12_rendered.docx")
    render_docx(ir, RESUME_TEMPLATE, out)
    return Document(out)


@pytest.mark.skipif(not _SAMPLE_12.exists(), reason="sample 12 PDF not in test fixtures")
class TestSample12LabelColTable:
    """CONTACT INFO section must render as a 3-col vMerge table, not standalone heading + 2-col table."""

    def test_three_col_vmerge_table_exists(self, tmp_path):
        """At least one table with 3 columns and w:vMerge restart in col 0 must exist."""
        d = _render_sample_12(tmp_path)
        found = False
        for tbl in d.element.body.iter(f"{{{_WM}}}tbl"):
            rows = tbl.findall(f"{{{_WM}}}tr")
            if not rows:
                continue
            first_row_cells = rows[0].findall(f"{{{_WM}}}tc")
            if len(first_row_cells) != 3:
                continue
            # Col 0 must have w:vMerge val="restart"
            tc0Pr = first_row_cells[0].find(f"{{{_WM}}}tcPr")
            if tc0Pr is None:
                continue
            vm = tc0Pr.find(f"{{{_WM}}}vMerge")
            if vm is not None and vm.get(f"{{{_WM}}}val") == "restart":
                found = True
                break
        assert found, (
            "Expected a 3-col table with w:vMerge restart in col 0 for CONTACT INFO label-column layout"
        )

    def test_label_col_h_rule_is_standalone_para(self, tmp_path):
        """An h_rule standalone paragraph (w:pBdr/top) must appear immediately before the 3-col vMerge table."""
        d = _render_sample_12(tmp_path)
        body_children = list(d.element.body)
        for i, child in enumerate(body_children):
            if child.tag != f"{{{_WM}}}tbl":
                continue
            rows = child.findall(f"{{{_WM}}}tr")
            if not rows:
                continue
            cells = rows[0].findall(f"{{{_WM}}}tc")
            if len(cells) != 3:
                continue
            tc0Pr = cells[0].find(f"{{{_WM}}}tcPr")
            if tc0Pr is None:
                continue
            vm = tc0Pr.find(f"{{{_WM}}}vMerge")
            if vm is None or vm.get(f"{{{_WM}}}val") != "restart":
                continue
            # Found the label-col table — check the element immediately before it
            if i == 0:
                break
            prev = body_children[i - 1]
            if prev.tag != f"{{{_WM}}}p":
                break
            pPr = prev.find(f"{{{_WM}}}pPr")
            if pPr is None:
                break
            pBdr = pPr.find(f"{{{_WM}}}pBdr")
            if pBdr is None:
                break
            top = pBdr.find(f"{{{_WM}}}top")
            if top is not None and top.get(f"{{{_WM}}}val") == "single":
                return  # test passes
            break
        pytest.fail("Expected a w:pBdr/top standalone paragraph immediately before the 3-col label-col table")

    def test_header_para_color_preserved(self, tmp_path):
        """Header paragraphs (name, title) must carry the template accent color from the PDF."""
        from tailor.compiler.pdf_parser import parse_pdf

        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        # All header paras should have a non-None text_color (the dark red #943613)
        header_colors = [
            pm.paragraph_profile.text_color
            for pm in ir.header_paras
            if pm.paragraph_profile
        ]
        assert any(c is not None for c in header_colors), (
            f"Expected at least one header_para to have a non-None text_color, got: {header_colors}"
        )

    def test_header_para_color_theme_override(self, tmp_path):
        """Rendered DOCX runs with explicit text_color must set w:themeColor='none' so
        LibreOffice uses the direct val instead of the paragraph style's theme color."""
        from tailor.compiler.pdf_parser import parse_pdf
        from tailor.compiler.docx_renderer import render_docx
        from tailor.config import RESUME_TEMPLATE
        from docx import Document

        _WN = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

        with open(_SAMPLE_12, "rb") as f:
            ir = parse_pdf(f.read())

        out = str(tmp_path / "s12_color_check.docx")
        render_docx(ir, RESUME_TEMPLATE, out)
        d = Document(out)

        colored_runs = []
        for p in d.element.body.iter(f"{{{_WN}}}p"):
            for r in p.findall(f"{{{_WN}}}r"):
                rPr = r.find(f"{{{_WN}}}rPr")
                if rPr is None:
                    continue
                color = rPr.find(f"{{{_WN}}}color")
                if color is None:
                    continue
                val = color.get(f"{{{_WN}}}val")
                if val and val not in ("000000", "FFFFFF", "auto"):
                    colored_runs.append((val, color.get(f"{{{_WN}}}themeColor")))

        assert colored_runs, "Expected at least one colored run in the rendered DOCX"
        for val, tc in colored_runs:
            assert tc == "none", (
                f"Run with explicit color val={val} must have themeColor='none', got {tc!r}"
            )


@pytest.mark.skipif(not _SAMPLE_12.exists(), reason="sample 12 PDF not in test fixtures")
class TestSample12GroupTableTopBorder:
    """EDU/COMM/LEAD group table must have a standalone w:pBdr/top paragraph immediately before it."""

    def _find_group_table(self, d):
        """Return (index, tbl) of the EDU/COMM/LEAD group table (N>=2 cols, no vMerge restart)."""
        body_children = list(d.element.body)
        for i, child in enumerate(body_children):
            if child.tag != f"{{{_WM}}}tbl":
                continue
            rows = child.findall(f"{{{_WM}}}tr")
            if not rows:
                continue
            cells = rows[0].findall(f"{{{_WM}}}tc")
            if len(cells) < 2:
                continue
            tc0Pr = cells[0].find(f"{{{_WM}}}tcPr")
            if tc0Pr is not None:
                vm = tc0Pr.find(f"{{{_WM}}}vMerge")
                if vm is not None and vm.get(f"{{{_WM}}}val") == "restart":
                    continue
            return i, child
        return None, None

    def test_group_table_preceded_by_h_rule_para(self, tmp_path):
        """A standalone w:pBdr/top paragraph must immediately precede the EDU/COMM/LEAD group table."""
        d = _render_sample_12(tmp_path)
        body_children = list(d.element.body)
        idx, tbl = self._find_group_table(d)
        assert tbl is not None, "EDU/COMM/LEAD group table not found in rendered DOCX"
        assert idx > 0, "Group table is the first body element — no preceding paragraph"
        prev = body_children[idx - 1]
        assert prev.tag == f"{{{_WM}}}p", (
            f"Element before group table should be a paragraph, got: {prev.tag}"
        )
        pPr = prev.find(f"{{{_WM}}}pPr")
        assert pPr is not None, "Preceding paragraph has no pPr"
        pBdr = pPr.find(f"{{{_WM}}}pBdr")
        assert pBdr is not None, "Preceding paragraph has no pBdr — h_rule paragraph missing"
        top = pBdr.find(f"{{{_WM}}}top")
        assert top is not None and top.get(f"{{{_WM}}}val") == "single", (
            "Expected w:pBdr/top val='single' on the h_rule paragraph before the group table"
        )

    def test_education_heading_has_no_para_border(self, tmp_path):
        """EDUCATION section heading inside the group table cell must NOT have w:pBdr/top (it moved to the table)."""
        d = _render_sample_12(tmp_path)
        # Find the group table (N>=2 cols, no vMerge, has tblBorders/top single)
        for tbl in d.element.body.iter(f"{{{_WM}}}tbl"):
            rows = tbl.findall(f"{{{_WM}}}tr")
            if not rows:
                continue
            first_row_cells = rows[0].findall(f"{{{_WM}}}tc")
            if len(first_row_cells) < 2:
                continue
            tc0Pr = first_row_cells[0].find(f"{{{_WM}}}tcPr")
            if tc0Pr is not None:
                vm = tc0Pr.find(f"{{{_WM}}}vMerge")
                if vm is not None and vm.get(f"{{{_WM}}}val") == "restart":
                    continue
            tblPr = tbl.find(f"{{{_WM}}}tblPr")
            if tblPr is None:
                continue
            tblBorders = tblPr.find(f"{{{_WM}}}tblBorders")
            if tblBorders is None:
                continue
            top = tblBorders.find(f"{{{_WM}}}top")
            if top is None or top.get(f"{{{_WM}}}val") != "single":
                continue
            # Found the group table — check that no cell paragraph has w:pBdr/top
            for p in tbl.iter(f"{{{_WM}}}p"):
                pPr = p.find(f"{{{_WM}}}pPr")
                if pPr is None:
                    continue
                pBdr = pPr.find(f"{{{_WM}}}pBdr")
                if pBdr is None:
                    continue
                p_top = pBdr.find(f"{{{_WM}}}top")
                if p_top is not None and p_top.get(f"{{{_WM}}}val") not in ("none", "nil", None):
                    assert False, (
                        "Group table cell heading must not have w:pBdr/top — h_rule moved to w:tblBorders/top"
                    )
            return  # Found and verified the group table
        pytest.skip("Group table with top border not found — test is conditional on layout detection")
