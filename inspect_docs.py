import sys
import io

# Force UTF-8 output to avoid cp1252 encoding errors on Windows console
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document

print('=== TEMPLATE ===')
doc = Document('templates/Leonid_Verman_Resume_Template.docx')
for i, p in enumerate(doc.paragraphs):
    is_bold = any(r.bold for r in p.runs)
    has_hyperlink = 'hyperlink' in p._p.xml
    text = p.text[:80]
    print(f'[{i}] style={repr(p.style.name):20} bold={is_bold} hyp={has_hyperlink} text={repr(text)}')

print()
print('=== OUTPUT ===')
doc2 = Document('output/Promise Robotics_Resume.docx')
for i, p in enumerate(doc2.paragraphs):
    is_bold = any(r.bold for r in p.runs)
    has_pipe = '|' in p.text
    text = p.text[:80]
    print(f'[{i}] style={repr(p.style.name):20} bold={is_bold} pipe={has_pipe} text={repr(text)}')
