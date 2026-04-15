"""
backend/app/services/resume_parser_service.py

Resume parser service — converts uploaded DOCX/PDF to a
StructuredResumeDocument stored in the DB.

Parsing strategy
----------------
1. Extract raw text from the uploaded file (DOCX via python-docx,
   PDF via pypdf).
2. Run lightweight heuristic parsing to populate StructuredResumeDocument
   fields where possible.
3. Preserve the raw_text field for downstream use (e.g. feeding the
   generator's prompt directly, which currently uses the template DOCX
   text anyway).

The heuristic parser is intentionally minimal — a full NLP-based parser
is a future enhancement.  The important thing is that the DB record
exists and the raw text is available.
"""

from __future__ import annotations

import io
import logging
import re

from backend.app.schemas.structured_resume import (
    ContactInfo,
    ExperienceEntry,
    SkillsSection,
    StructuredResumeDocument,
)

logger = logging.getLogger(__name__)


# ── Text extraction ────────────────────────────────────────────────────────

def extract_text_from_docx(data: bytes) -> str:
    """Extract plain text from a DOCX file given as bytes.

    Handles both standard DOCX files and LibreOffice-converted files.

    LibreOffice PDF→DOCX conversion wraps every text box inside an
    ``<mc:AlternateContent>`` element that contains both a ``<mc:Choice>``
    (real content) and a ``<mc:Fallback>`` (duplicate for old viewers).
    Standard ``doc.paragraphs`` only sees body-level paragraphs and misses
    text boxes entirely, so we collect paragraphs from the full XML tree
    while skipping anything inside a ``<mc:Fallback>`` to avoid duplicates.
    """
    from docx import Document
    from docx.oxml.ns import qn

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        from zipfile import BadZipFile
        from fastapi import HTTPException, status
        if isinstance(exc, BadZipFile) or "not a zip file" in str(exc).lower():
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=(
                    "The file could not be read as a DOCX document. "
                    "It may be in the older .doc format or be corrupted. "
                    "Please save it as .docx or convert it to PDF and try again."
                ),
            ) from exc
        raise

    W  = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    _FALLBACK = "{%s}Fallback" % MC
    _T        = "{%s}t" % W
    _P        = "{%s}p" % W

    body = doc.element.body

    def _in_fallback(elem) -> bool:
        """Return True if elem is a descendant of an mc:Fallback element.

        lxml proxy objects are recreated on each iter() call so id()-based
        sets are unreliable.  Walking the parent chain is O(depth) but
        always correct.
        """
        parent = elem.getparent()
        while parent is not None:
            if parent.tag == _FALLBACK:
                return True
            parent = parent.getparent()
        return False

    # Walk all paragraphs in document order, skipping Fallback descendants.
    lines: list[str] = []
    for para in body.iter(_P):
        if _in_fallback(para):
            continue
        text = "".join(
            elem.text
            for elem in para.iter(_T)
            if elem.text and not _in_fallback(elem)
        )
        if text.strip():
            lines.append(text)

    return "\n".join(lines)


def extract_text_from_pdf(data: bytes) -> str:
    """Extract plain text from a PDF file given as bytes."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def extract_text(data: bytes, filename: str) -> str:
    """Dispatch to the correct extractor based on file extension."""
    name_lower = filename.lower()
    if name_lower.endswith(".docx"):
        return extract_text_from_docx(data)
    if name_lower.endswith(".pdf"):
        return extract_text_from_pdf(data)
    # Try DOCX first, fall back to treating as text
    try:
        return extract_text_from_docx(data)
    except Exception:
        return data.decode("utf-8", errors="replace")


# ── Heuristic structure extraction ────────────────────────────────────────

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[a-zA-Z]{2,}")
_PHONE_RE = re.compile(r"\+?[\d\s\-().]{7,20}")
_LINKEDIN_RE = re.compile(r"linkedin\.com/in/[\w-]+", re.IGNORECASE)
_GITHUB_RE = re.compile(r"github\.com/[\w-]+", re.IGNORECASE)


def _extract_contacts(text: str) -> ContactInfo:
    email = next(iter(_EMAIL_RE.findall(text)), None)
    phone_matches = _PHONE_RE.findall(text[:500])  # only search header area
    phone = phone_matches[0].strip() if phone_matches else None
    linkedin = next(iter(_LINKEDIN_RE.findall(text)), None)
    linkedin_url = f"https://{linkedin}" if linkedin else None
    github = next(iter(_GITHUB_RE.findall(text)), None)
    github_url = f"https://{github}" if github else None
    return ContactInfo(
        email=email,
        phone=phone,
        linkedin_url=linkedin_url,
        github_url=github_url,
    )


def _extract_name(lines: list[str]) -> str:
    """Best-effort: return the candidate name from the first promising line.

    Strategy: strip known non-name tokens (email, phone, social URLs) from
    each line, then collect the leading run of title-cased word tokens.
    Require at least two such tokens (first + last name) to avoid returning
    section headers like "Experience".

    Example: "Gleb Zernov ć gzernov@proton.me"
        → strip email → "Gleb Zernov ć"
        → leading title-cased tokens → ["Gleb", "Zernov"]
        → return "Gleb Zernov"
    """
    for line in lines[:8]:
        stripped = line.strip()
        if not stripped:
            continue
        # Remove non-name tokens so the name isn't discarded just because it
        # shares a line with a phone number, email, or social URL.
        candidate = _EMAIL_RE.sub("", stripped)
        candidate = _PHONE_RE.sub("", candidate)
        candidate = _LINKEDIN_RE.sub("", candidate)
        candidate = _GITHUB_RE.sub("", candidate)
        candidate = candidate.strip()
        if not candidate:
            continue
        # Collect the leading run of title-cased tokens (≥ 2 letters, starts
        # with an uppercase letter).  Stop at the first token that breaks the
        # pattern so garbled suffixes like "ć" don't end up in the name.
        name_tokens: list[str] = []
        for word in candidate.split():
            alpha_only = "".join(c for c in word if c.isalpha())
            if len(alpha_only) >= 2 and alpha_only[0].isupper():
                name_tokens.append(alpha_only)
            elif not name_tokens:
                continue  # skip leading non-name junk before the name starts
            else:
                break  # first non-name token after name started → stop
        if len(name_tokens) >= 2:
            return " ".join(name_tokens[:4])
    return "Unknown"


def _heuristic_parse(raw_text: str) -> StructuredResumeDocument:
    """Very lightweight structure extraction from raw resume text."""
    lines = raw_text.splitlines()
    non_empty = [l for l in lines if l.strip()]

    name = _extract_name(non_empty)
    contacts = _extract_contacts(raw_text)

    # Summary: look for a "summary" or "profile" section header
    summary = ""
    for i, line in enumerate(non_empty):
        if re.search(r"\b(summary|profile|objective)\b", line, re.IGNORECASE):
            # Collect next 1–3 non-empty lines as summary
            snippet_lines = [l.strip() for l in non_empty[i + 1 : i + 5] if l.strip()]
            summary = " ".join(snippet_lines[:3])
            break

    return StructuredResumeDocument(
        name=name,
        contacts=contacts,
        summary=summary,
        experience=[],          # heuristic experience extraction is future work
        technical_skills=SkillsSection(),
        raw_text=raw_text,
    )


# ── Service class ──────────────────────────────────────────────────────────

class ResumeParserService:
    """
    Parse an uploaded resume file into a StructuredResumeDocument.

    Responsibilities
    ----------------
    - extract text from DOCX or PDF
    - run heuristic structure extraction
    - return a StructuredResumeDocument suitable for DB storage

    The DB write itself is handled by the caller (API layer or generation
    service) so this service remains pure and testable.
    """

    def parse(
        self,
        data: bytes,
        filename: str = "resume.docx",
    ) -> StructuredResumeDocument:
        """
        Parse resume bytes into a StructuredResumeDocument.

        Parameters
        ----------
        data:
            Raw file bytes (DOCX or PDF).
        filename:
            Original filename; used to detect file type.
        """
        logger.debug("Parsing resume file=%s size=%d bytes", filename, len(data))
        raw_text = extract_text(data, filename)
        if not raw_text.strip():
            logger.warning("Resume file %s yielded empty text", filename)
            return StructuredResumeDocument(name="Unknown", raw_text="")

        doc = _heuristic_parse(raw_text)
        logger.debug("Parsed resume name=%r summary_len=%d", doc.name, len(doc.summary))
        return doc

    def parse_raw_text(self, text: str) -> StructuredResumeDocument:
        """Parse from already-extracted text (e.g. pasted in UI)."""
        return _heuristic_parse(text)
